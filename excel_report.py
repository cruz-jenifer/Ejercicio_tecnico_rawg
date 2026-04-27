import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def generar_reporte_excel(ruta_csv, ruta_excel):
    """lee el csv generado y crea un reporte en excel ordenado por requerimientos"""

    # lectura del csv
    try:
        with open(ruta_csv, mode='r', encoding='utf-8') as archivo_csv:
            lector = csv.DictReader(archivo_csv)
            datos = list(lector)
    except FileNotFoundError:
        print(f" Error: no se encontró el archivo CSV en '{ruta_csv}'")
        return
    except IOError as e:
        print(f"Error de lectura al abrir '{ruta_csv}': {e}")
        return

    if not datos:
        print("el archivo CSV está vacío, no se genera reporte")
        return

    # ordena por requerimientos mínimos (los N/A van al final)
    datos_ordenados = sorted(
        datos,
        key=lambda x: (
            x.get("Requerimientos mínimos", "N/A") == "N/A",
            x.get("Requerimientos mínimos", "")
        )
    )

    # creacion del libro excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Videojuegos"

    # cabeceras
    cabeceras = ["Nombre", "Rating", "Fecha de Lanzamiento", "Requerimientos Mínimos"]
    ws.append(cabeceras)

    # estilos de cabecera
    fuente_negrita = Font(bold=True, color="FFFFFF")
    fondo_color = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    for celda in ws[1]:
        celda.font = fuente_negrita
        celda.fill = fondo_color

    # inserta los datos ordenados
    for juego in datos_ordenados:
        fila = [
            juego.get("Nombre del videojuego", "Desconocido"),
            juego.get("Rating", "0"),
            juego.get("Fecha de lanzamiento", "N/A"),
            juego.get("Requerimientos mínimos", "N/A")
        ]
        ws.append(fila)

    # ajuste de ancho de columnas
    for columna in ws.columns:
        longitud_maxima = 0
        letra_columna = columna[0].column_letter

        for celda in columna:
            if celda.value is not None:
                longitud_celda = len(str(celda.value))
                if longitud_celda > longitud_maxima:
                    longitud_maxima = longitud_celda

        # margen extra para legibilidad
        ws.column_dimensions[letra_columna].width = longitud_maxima + 2

    # guardado del archivo
    try:
        wb.save(ruta_excel)
        print(f" Reporte generado exitosamente en: {ruta_excel}")
    except IOError:
        print(f" No se pudo guardar el Excel. verifica que '{ruta_excel}' no esté abierto")
